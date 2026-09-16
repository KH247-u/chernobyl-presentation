from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

def create_pptx():
    # Load HTML
    with open('index.html', 'r', encoding='utf-8') as f:
        soup = BeautifulSoup(f, 'html.parser')

    # Create Presentation
    prs = Presentation()
    
    # Define Layouts
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]
    
    slides = soup.find_all('section', class_='slide')
    
    for slide_idx, slide_soup in enumerate(slides):
        # Determine if Title Slide
        is_title = slide_idx == 0 or slide_idx == len(slides) - 1 or slide_soup.find(class_='main-title') is not None
        
        layout = title_slide_layout if is_title else bullet_slide_layout
        slide = prs.slides.add_slide(layout)
        
        shapes = slide.shapes
        
        title_shape = shapes.title
        body_shape = shapes.placeholders[1] if len(shapes.placeholders) > 1 else None
        
        # Extract Text Elements
        h1 = slide_soup.find('h1')
        h2 = slide_soup.find('h2')
        h3 = slide_soup.find('h3')
        
        title_text = ""
        if h1: title_text = h1.get_text(separator=' ').strip()
        elif h2: title_text = h2.get_text(separator=' ').strip()
        elif h3: title_text = h3.get_text(separator=' ').strip()
        
        if title_shape and title_text:
            title_shape.text = title_text
            
        # Extract body text (bullets, paragraphs, cards)
        if body_shape:
            tf = body_shape.text_frame
            tf.text = "" # Clear default
            
            # Subtitle or leading text
            subtitle = slide_soup.find(class_='subtitle')
            if subtitle:
                p = tf.add_paragraph()
                p.text = subtitle.get_text(separator=' ').strip()
                p.level = 0
            
            # List items
            ul = slide_soup.find('ul')
            if ul:
                for li in ul.find_all('li'):
                    p = tf.add_paragraph()
                    p.text = li.get_text(separator=' ').strip()
                    p.level = 0
                    
            # Other paragraphs if no UL
            if not ul:
                paragraphs = slide_soup.find_all('p')
                for para in paragraphs:
                    text = para.get_text(separator=' ').strip()
                    if text:
                        p = tf.add_paragraph()
                        p.text = text
                        p.level = 0
                        
            # Specific cards like team profiles, comparison cards
            team_profiles = slide_soup.find_all(class_='team-profile-card')
            for profile in team_profiles:
                name = profile.find(class_='profile-name')
                if name:
                    p = tf.add_paragraph()
                    p.text = "Team Member: " + name.get_text(separator=' ').strip()
                    p.level = 0
                    
            cards = slide_soup.find_all(class_=['lesson-card', 'concl-card', 'health-card', 'env-metric-card'])
            for card in cards:
                card_title = card.find(['h4', '.hc-title', '.metric-lbl'])
                card_desc = card.find('p')
                card_val = card.find(class_='metric-val')
                
                text = ""
                if card_title: text += card_title.get_text(separator=' ').strip() + ": "
                if card_val: text += card_val.get_text(separator=' ').strip() + " - "
                if card_desc: text += card_desc.get_text(separator=' ').strip()
                
                if text:
                    p = tf.add_paragraph()
                    p.text = text
                    p.level = 0

    pptx_path = 'Chernobyl_Disaster_Presentation.pptx'
    prs.save(pptx_path)
    print(f"Successfully generated {pptx_path}")

if __name__ == '__main__':
    create_pptx()
